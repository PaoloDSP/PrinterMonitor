"""Moteur de DIAGNOSTIC du rapprochement de parc.

Exécute le moteur complet sur les fichiers réels et produit un rapport
texte : correspondances, traductions inconnues/ambiguës, vrais changements,
sans changement, ajouts, retraits. Aucune écriture dans la CMDB.
"""

import os

from . import conventions
from .importer import load_client_file
from .reconcile import ReconcileEngine
from .reconcile import CmdbRow


def load_cmdb_rows(path):
    """Charge les lignes CMDB depuis la feuille « CMDB » d'un classeur."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "CMDB" in wb.sheetnames:
            ws = wb["CMDB"]
        elif len(wb.sheetnames) >= 3:
            ws = wb[wb.sheetnames[2]]
        else:
            ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    header_row = None
    for idx, row in enumerate(rows):
        if row and str(row[0]).strip() == "Grandsite":
            header_row = idx
            break
    if header_row is None:
        raise ValueError("En-têtes CMDB non reconnus (colonne « Grandsite »).")
    out = []
    for row_idx, row in enumerate(rows[header_row + 1:], start=header_row + 2):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        out.append(CmdbRow(row_idx, row))
    return out


def run_diagnostic(client_path, cmdb_path, mappings=None, verbose=False):
    """Exécute le rapprochement complet. Renvoie (result, engine)."""
    devices, sites_rms, _info = load_client_file(client_path)
    cmdb_rows = load_cmdb_rows(cmdb_path)
    engine = ReconcileEngine(cmdb_rows, devices, sites_rms=sites_rms,
                             mappings=mappings)
    result = engine.run()
    return result, engine


def format_report(result):
    """Rapport lisible du résultat du rapprochement."""
    d = result.diagnostics
    lines = []
    add = lines.append
    add("=" * 72)
    add("RAPPORT DE RAPPROCHEMENT DE PARC")
    add("=" * 72)
    add(f"Équipements analysés          : {d.get('equipements_analyses', 0)}")
    add(f"Lignes CMDB                   : {d.get('lignes_cmdb', 0)}")
    add("-" * 72)
    add(f"Correspondances exactes       : {d.get('correspondances_exactes', 0)}"
        "   (tatouage)")
    add(f"Correspondances historiques   : {d.get('correspondances_historiques', 0)}"
        "   (n° série / MAC)")
    add(f"Correspondances ambiguës      : {d.get('correspondances_ambigues', 0)}")
    add("-" * 72)
    add("Traductions inconnues (anomalies) :")
    add(f"   modèles inconnus            : {d.get('modeles_inconnus', 0)}")
    add(f"   sites inconnus              : {d.get('sites_inconnus', 0)}")
    add(f"   bâtiments inconnus          : {d.get('batiments_inconnus', 0)}")
    add(f"   étages inconnus             : {d.get('etages_inconnus', 0)}")
    add(f"   bureaux inconnus            : {d.get('bureaux_inconnus', 0)}")
    add("-" * 72)
    add(f"Vrais changements             : {d.get('vraies_modifications', 0)}")
    add(f"Sans changement               : {d.get('sans_changement', 0)}")
    add(f"Ajouts                        : {d.get('ajouts', 0)}")
    add(f"Retraits potentiels           : {d.get('retraits_potentiels', 0)}")
    add("Retraits écartés (états métier): "
        f"{d.get('retraits_ignores_etat_special', 0)}")
    add("Conventions appliquées        : "
        f"{d.get('conventions_appliquees', 0)}"
        "   (valeurs équivalentes, aucune modification)")
    add("-" * 72)

    # États métier : le matériel est connu, mais son adresse n'est pas une
    # donnée exploitable. Ces lignes ne sont ni des retraits ni des ajouts :
    # elles doivent rester lisibles plutôt que d'être noyées (§3, §4).
    states = d.get("etats_speciaux_cmdb") or {}
    if states:
        add("Lignes CMDB en état métier (jamais proposées en retrait) :")
        for state, n in sorted(states.items(), key=lambda kv: (-kv[1], kv[0])):
            add(f"   {conventions.state_label(state):<45}: {n:6d}")
        add(f"   anomalies d'état non comparables (équipements appariés) : "
            f"{d.get('anomalies_etat_special', 0)}")
        add(f"   réapparitions (CMDB « retiré », présence client)        : "
            f"{d.get('anomalies_reapparition', 0)}")
        add("-" * 72)

    levels = d.get("niveaux_confiance", {})
    total_fields = sum(levels.values())
    if total_fields:
        add("Niveaux de confiance des champs traduits :")
        for label in ("Exact", "Sites RMS", "Historique (CMDB)",
                      "Normalisé", "Flou", "Inconnu"):
            n = levels.get(label, 0)
            pct = 100.0 * n / total_fields
            add(f"   {label:<20}: {n:6d}  ({pct:.1f} %)")
    add("-" * 72)

    groups = {}
    for anomaly in result.anomalies:
        groups.setdefault(anomaly.kind, []).append(anomaly)
    add("Anomalies groupées :")
    for kind in sorted(groups):
        items = sorted(groups[kind], key=lambda a: -a.count)
        add(f"   [{kind}] {len(items)} groupe(s), "
            f"{sum(a.count for a in items)} équipement(s)")
        for anomaly in items[:12]:
            line = f"      • {anomaly.label}"
            if anomaly.candidates:
                cands = " | ".join(anomaly.candidates[:3])
                line += f"   -> {cands}"
            line += f"   ({anomaly.count} éq.)"
            add(line)
        if len(items) > 12:
            add(f"      … et {len(items) - 12} autres groupes")
    add("=" * 72)

    # Conventions appliquées : ce que la normalisation a absorbé, groupe par
    # groupe. Un nommage différent n'est PAS une modification (§23) : il vaut
    # mieux lire « CHOOZ AMEN → CHOOZ : 154 équipements, aucune modification »
    # que 154 anomalies.
    lines_conventions = conventions.describe_conventions(
        result.conventions, limit=20)
    if lines_conventions:
        add("Conventions appliquées (différences de nommage/formatage"
            " — aucune modification) :")
        lines.extend("   • " + line for line in lines_conventions)
        rest = sum(result.conventions.values()) - sum(
            n for (_f, _s, _d, _k, _l), n in sorted(
                result.conventions.items(), key=lambda kv: (-kv[1], kv[0])
            )[:20]
        )
        if rest > 0:
            add(f"   … et {rest} autre(s) application(s) de convention")
        add("=" * 72)
    return "\n".join(lines)


if __name__ == "__main__":
    if __package__ in (None, ""):
        # Exécution directe : remonte au répertoire du projet.
        import sys
        sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(
            os.path.abspath(__file__)))))
        from printer_monitor.presta.diagnostic import (
            run_diagnostic, format_report)
    base = os.path.dirname(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))))
    client = os.path.join(base, "Presta RMS.xlsm")
    cmdb = None
    for name in os.listdir(base):
        if name.startswith("RMS_Gestion_consommables") and \
                name.endswith(".xlsx"):
            cmdb = os.path.join(base, name)
            break
    if not (os.path.isfile(client) and cmdb):
        raise SystemExit("Fichiers réels introuvables pour le diagnostic.")
    result, _engine = run_diagnostic(client, cmdb)
    print(format_report(result))